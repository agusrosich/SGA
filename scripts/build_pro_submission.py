"""Build the Practical Radiation Oncology (PRO) submission package in envio_PRO/.

Usage (from the repository root):
    python scripts/build_pro_submission.py

Reads manuscrito_final/manuscript.md (single source of truth for the text) and
the part 1 outputs, and writes:
    01_Cover_Letter.docx, 02_Title_Page.docx, 03_Manuscript_Anonymized.docx,
    04_Figure_Legends.docx, 05_Supplementary_Material.docx, Figures/Figure_N.tif

PRO requires an anonymized manuscript and a separate title page, so author
names and affiliations only appear in the title page and cover letter.
Placeholders the authors must fill are written as [COMPLETAR: ...] and
highlighted in yellow.
"""

from __future__ import annotations

import importlib.util
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import common.causal_core as core

MANUSCRIPT = ROOT / "manuscrito_final" / "manuscript.md"
PART1 = ROOT / "01_entrenamiento_seer"
OUT = ROOT / "envio_PRO"

RUNNING_TITLE = "Adding chemotherapy to RT in salivary cancer"
AUTHORS = [
    ("Federico Lorenzo", "1,2"), ("Agustin Rosich", "1,2"), ("Jesica Lell", "1,2"),
    ("Sergio Aguiar", "1"), ("Valentina Ferreira", "1"), ("Karina Ochandorena", "1,2"),
    ("Eduardo Larrinaga", "1"), ("Natalia Gadea", "1"), ("Nicolas Larragueta", "1"),
    ("Aldo Quarneti", "1"),
]
AFFILIATIONS = [
    ("1", "Radiotherapy, RT International Institute, Montevideo, Uruguay."),
    ("2", "Radiotherapy, Unidad Academica de Radioterapia, Montevideo, Uruguay, RT International Institute."),
]
COI_STATEMENT = (
    "None. All authors declare that they have no known competing financial interests or personal "
    "relationships that could have appeared to influence the work reported in this paper."
)
FUNDING_STATEMENT = (
    "[COMPLETAR: fuente de financiamiento, o bien \"This research did not receive any specific grant from "
    "funding agencies in the public, commercial, or not-for-profit sectors.\"]"
)
DATA_SHARING_STATEMENT = (
    "The data analyzed in this study were obtained from the Surveillance, Epidemiology, and End Results "
    "(SEER) Program of the National Cancer Institute and are available to investigators through the SEER "
    "Program upon completion of its data-use agreement (https://seer.cancer.gov/data/). The analysis code "
    "is available from the corresponding author upon reasonable request."
)
PLACEHOLDER = re.compile(r"(\[COMPLETAR:[^\]]*\])")
SUPERSCRIPT = re.compile(r"(\^[0-9,]+)")


# ---------------------------------------------------------------- markdown model

@dataclass
class Table:
    number: str
    caption: str
    header: list[str]
    rows: list[list[str]]
    right_aligned: list[bool]
    footnote: str = ""


@dataclass
class Figure:
    number: str
    caption: str
    path: Path


@dataclass
class Manuscript:
    title: str = ""
    abstract: list[tuple[str, str]] = field(default_factory=list)
    body: list[tuple[str, str]] = field(default_factory=list)  # ("h2"|"h3"|"p", text)
    tables: list[Table] = field(default_factory=list)
    figures: list[Figure] = field(default_factory=list)
    references: list[str] = field(default_factory=list)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_manuscript(path: Path) -> Manuscript:
    lines = path.read_text(encoding="utf-8").splitlines()
    m = Manuscript()
    section = None
    pending_caption: tuple[str, str] | None = None
    last_table: Table | None = None
    i = 0
    while i < len(lines):
        s = lines[i].strip(); i += 1
        if not s or s.startswith("Authors:"):
            continue
        if s.startswith("# "):
            m.title = s[2:].strip(); continue
        h = re.match(r"^(#{2,3})\s+(.*)$", s)
        if h:
            level, text = len(h.group(1)), h.group(2).strip()
            if level == 2:
                section = text
                if text not in ("Abstract", "References"):
                    m.body.append(("h2", text))
            elif section == "Abstract":
                m.abstract.append((text, ""))
            else:
                m.body.append(("h3", text))
            last_table = None
            continue
        if section == "References":
            ref = re.match(r"^\d+\.\s+(.*)$", s)
            if ref:
                m.references.append(ref.group(1))
            continue
        fig = re.match(r"^!\[Figure (\d+)\.\s*(.*)\]\(([^)]+)\)$", s)
        if fig:
            m.figures.append(Figure(fig.group(1), fig.group(2), path.parent / fig.group(3)))
            last_table = None
            continue
        cap = re.match(r"^Table (\d+)\.\s*(.*)$", s)
        if cap:
            pending_caption = (cap.group(1), cap.group(2)); continue
        if s.startswith("|"):
            block = [s]
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            if pending_caption is None:
                raise ValueError(f"Table without a numbered caption near: {s[:60]}")
            aligns = split_row(block[1])
            last_table = Table(pending_caption[0], pending_caption[1], split_row(block[0]),
                               [split_row(r) for r in block[2:]], [a.endswith(":") for a in aligns])
            m.tables.append(last_table); pending_caption = None
            continue
        if last_table is not None and not last_table.footnote:
            last_table.footnote = s; last_table = None; continue
        if section == "Abstract":
            label, _ = m.abstract[-1]; m.abstract[-1] = (label, s)
        else:
            m.body.append(("p", s))
    return m


def words(text: str) -> int:
    return len(re.findall(r"\S+", text))


# ---------------------------------------------------------------- docx helpers

def new_document(double_spaced: bool, line_numbers: bool) -> Document:
    doc = Document()
    props = doc.core_properties
    props.author = ""; props.last_modified_by = ""; props.comments = ""; props.keywords = ""
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2.0 if double_spaced else 1.15
    normal.paragraph_format.space_after = Pt(0 if double_spaced else 6)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
        st.font.italic = name == "Heading 3"; st.font.color.rgb = RGBColor(0, 0, 0)
        # Theme font attributes override the explicit font name, so drop them.
        rfonts = st.element.rPr.find(qn("w:rFonts"))
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts is not None and rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    if line_numbers:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
        sec._sectPr.append(ln)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; run.append(t); fld.append(run)
    footer._p.append(fld)
    return doc


def add_runs(paragraph, text: str, bold: bool = False, size: Pt | None = None) -> None:
    """Add text, highlighting [COMPLETAR: ...] placeholders and rendering ^1,2 as superscript."""
    for piece in PLACEHOLDER.split(text):
        if not piece:
            continue
        if PLACEHOLDER.fullmatch(piece):
            r = paragraph.add_run(piece); r.bold = bold
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if size: r.font.size = size
            continue
        for sub in SUPERSCRIPT.split(piece):
            if not sub:
                continue
            r = paragraph.add_run(sub[1:] if sub.startswith("^") else sub)
            r.bold = bold; r.font.superscript = sub.startswith("^")
            if size: r.font.size = size


def labeled(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    add_runs(p, f"{label}: ", bold=True); add_runs(p, text)


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table(doc: Document, table: Table, label: str) -> None:
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    add_runs(p, f"{label}. ", bold=True); add_runs(p, table.caption)
    t = doc.add_table(rows=1, cols=len(table.header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, values in enumerate([table.header] + table.rows):
        cells = t.rows[0].cells if row_idx == 0 else t.add_row().cells
        for col, value in enumerate(values[: len(cells)]):
            para = cells[col].paragraphs[0]
            para.paragraph_format.line_spacing = 1.0; para.paragraph_format.space_after = Pt(0)
            if table.right_aligned[col] and col > 0:
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_runs(para, value, bold=row_idx == 0, size=Pt(10))
    if table.footnote:
        f = doc.add_paragraph(); f.paragraph_format.line_spacing = 1.0; f.paragraph_format.space_before = Pt(6)
        add_runs(f, table.footnote, size=Pt(10))


# ---------------------------------------------------------------- outputs

def build_figures(m: Manuscript, fig_dir: Path) -> list[Path]:
    """Re-render the part 1 figures at 600 dpi and save them as RGB LZW TIFF named by figure number."""
    spec = importlib.util.spec_from_file_location("part1", PART1 / "train_seer_model.py")
    part1 = importlib.util.module_from_spec(spec); spec.loader.exec_module(part1)
    d, _ = part1.load_data(PART1 / "data" / "raw" / "ExportadaSEER_Estandarizada.csv")
    e = core.estimate(d)
    tables = PART1 / "outputs" / "tables"
    saved = pd.read_csv(tables / "adjusted_survival_curves.csv")
    if not (saved.sort_values(["Treatment", "Month"]).Survival.values.round(10)
            == e["curves"].sort_values(["Treatment", "Month"]).Survival.values.round(10)).all():
        raise RuntimeError("Re-estimated survival curves differ from the saved part 1 outputs.")
    render = fig_dir / "_render"
    core.plots(render, e, pd.read_csv(tables / "bootstrap.csv"), pd.read_csv(tables / "balance.csv"), dpi=600)
    paths = []
    for f in m.figures:
        target = fig_dir / f"Figure_{f.number}.tif"
        with Image.open(render / f"{f.path.stem}.png") as im:
            im.convert("RGB").save(target, compression="tiff_lzw", dpi=(600, 600))
        paths.append(target)
    for png in render.glob("*.png"):
        png.unlink()
    render.rmdir()
    return paths


def build_manuscript(m: Manuscript, path: Path) -> None:
    doc = new_document(double_spaced=True, line_numbers=True)
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(t, m.title, bold=True)
    doc.add_heading("Abstract", level=1)
    for label, text in m.abstract:
        labeled(doc, label, text)
    page_break(doc)
    for kind, text in m.body:
        if kind == "h2":
            doc.add_heading(text, level=1)
        elif kind == "h3":
            doc.add_heading(text, level=2)
        else:
            add_runs(doc.add_paragraph(), text)
    page_break(doc)
    doc.add_heading("References", level=1)
    for n, ref in enumerate(m.references, 1):
        add_runs(doc.add_paragraph(), f"{n}. {ref}")
    for table in m.tables:
        page_break(doc); add_table(doc, table, f"Table {table.number}")
    page_break(doc)
    doc.add_heading("Figure Legends", level=1)
    for f in m.figures:
        p = doc.add_paragraph(); add_runs(p, f"Figure {f.number}. ", bold=True); add_runs(p, f.caption)
    doc.core_properties.title = m.title
    doc.save(path)


def build_title_page(m: Manuscript, path: Path, counts: dict) -> None:
    doc = new_document(double_spaced=False, line_numbers=False)
    labeled(doc, "Article type", "Original Report")
    p = doc.add_paragraph(); add_runs(p, "Title: ", bold=True); add_runs(p, m.title, bold=False)
    labeled(doc, "Running title", RUNNING_TITLE)
    p = doc.add_paragraph(); add_runs(p, "Authors: ", bold=True)
    add_runs(p, ", ".join(f"{name}^{aff}" for name, aff in AUTHORS))
    p = doc.add_paragraph(); add_runs(p, "Affiliations:", bold=True)
    for num, text in AFFILIATIONS:
        add_runs(doc.add_paragraph(), f"^{num} {text}")
    labeled(doc, "Corresponding author",
            "[COMPLETAR: nombre completo, dirección postal, teléfono y e-mail del autor de correspondencia]")
    labeled(doc, "Author responsible for statistical analysis",
            "[COMPLETAR: nombre completo, dirección postal, teléfono y e-mail]")
    labeled(doc, "Conflict of Interest Statement", COI_STATEMENT)
    labeled(doc, "Funding Statement", FUNDING_STATEMENT)
    labeled(doc, "Data Sharing Statement", DATA_SHARING_STATEMENT)
    labeled(doc, "Acknowledgments", "[COMPLETAR: agradecimientos, o bien \"None.\"]")
    labeled(doc, "Word count", (
        f"text {counts['text_words']:,} words; abstract {counts['abstract_words']} words; "
        f"{counts['tables']} tables; {counts['figures']} figures; {counts['references']} references; "
        "1 supplementary table"))
    doc.core_properties.title = "Title page"
    doc.save(path)


def build_legends(m: Manuscript, path: Path) -> None:
    doc = new_document(double_spaced=True, line_numbers=False)
    doc.add_heading("Figure Legends", level=1)
    for f in m.figures:
        p = doc.add_paragraph(); add_runs(p, f"Figure {f.number}. ", bold=True); add_runs(p, f.caption)
    doc.core_properties.title = "Figure legends"
    doc.save(path)


def build_supplement(m: Manuscript, path: Path) -> None:
    flow = pd.read_csv(PART1 / "outputs" / "tables" / "cohort_flow.csv")
    table = Table("S1", "Cohort selection from the standardized SEER-derived extract.",
                  ["Selection step", "Remaining, n", "Excluded at step, n"],
                  [[r.Step, f"{r.Remaining_N:,}", f"{r.Excluded_at_step_N:,}"] for r in flow.itertuples()],
                  [False, True, True],
                  "Criteria were applied sequentially. Surgery, postoperative intent of radiotherapy, M stage, "
                  "and diagnosis year are not available in the extract and could not be used as eligibility "
                  "criteria. SEER, Surveillance, Epidemiology, and End Results.")
    doc = new_document(double_spaced=False, line_numbers=False)
    doc.add_heading("Supplementary Material", level=1)
    add_runs(doc.add_paragraph(), m.title)
    add_table(doc, table, "Supplementary Table S1")
    doc.core_properties.title = "Supplementary material"
    doc.save(path)


def signed(value: float, plus: bool = False) -> str:
    """One-decimal number with a typographic minus, matching the manuscript."""
    return f"{value:{'+' if plus else ''}.1f}".replace("-", "−")


def build_cover_letter(m: Manuscript, path: Path, primary: pd.Series) -> None:
    doc = new_document(double_spaced=False, line_numbers=False)
    paragraphs = [
        "[COMPLETAR: fecha]",
        "Dear Editor,",
        f"We are pleased to submit our manuscript entitled “{m.title}” for consideration as an "
        "Original Report in Practical Radiation Oncology.",
        "Whether chemotherapy adds a survival benefit to postoperative radiotherapy in high-risk salivary "
        "gland cancer remains unresolved, and RTOG 1008 was designed to answer this question. Before its "
        f"results are available, we analyzed {int(primary.N):,} patients with T3/T4 disease from a "
        "SEER-derived cohort using censoring-aware causal survival methods (propensity-score weighting, a "
        "covariate-adjusted weighted Cox model, and standardized restricted mean survival time) to generate "
        "a falsifiable pre-results prediction of the average treatment effect. The adjusted 10-year "
        f"restricted mean survival time difference was {signed(primary.Delta_RMST_Months)} months (95% "
        f"bootstrap CI, {signed(primary.CI95_Lower)} to {signed(primary.CI95_Upper, plus=True)}) for radiotherapy plus "
        "chemotherapy versus radiotherapy alone, and the direction of the estimate persisted in "
        "histology-restricted sensitivity analyses.",
        "We believe this work is of interest to the readers of Practical Radiation Oncology because it "
        "addresses a frequent decision in routine practice and provides an explicit, testable benchmark "
        "against which the forthcoming randomized evidence can be interpreted.",
        "This manuscript is original, has not been published previously, and is not under consideration "
        "for publication elsewhere. All authors have read and approved the manuscript and declare no "
        "conflicts of interest. The study used deidentified, publicly available SEER data.",
        "Thank you for your consideration.",
        "Sincerely,",
        "[COMPLETAR: nombre y título del autor de correspondencia]\nOn behalf of all authors",
    ]
    for text in paragraphs:
        add_runs(doc.add_paragraph(), text)
    doc.core_properties.title = "Cover letter"
    doc.save(path)


def main() -> None:
    m = parse_manuscript(MANUSCRIPT)
    OUT.mkdir(exist_ok=True); fig_dir = OUT / "Figures"; fig_dir.mkdir(exist_ok=True)
    text_words = sum(words(t) for kind, t in m.body if kind == "p")
    counts = {
        "text_words": text_words,
        "abstract_words": sum(words(t) for _, t in m.abstract),
        "tables": len(m.tables), "figures": len(m.figures), "references": len(m.references),
    }
    primary = pd.read_csv(PART1 / "outputs" / "tables" / "primary_result.csv").iloc[0]
    build_cover_letter(m, OUT / "01_Cover_Letter.docx", primary)
    build_title_page(m, OUT / "02_Title_Page.docx", counts)
    build_manuscript(m, OUT / "03_Manuscript_Anonymized.docx")
    build_legends(m, OUT / "04_Figure_Legends.docx")
    build_supplement(m, OUT / "05_Supplementary_Material.docx")
    figures = build_figures(m, fig_dir)
    print("Counts:", counts, "| running title chars:", len(RUNNING_TITLE))
    print("Figures:", ", ".join(p.name for p in figures))
    print(f"Output: {OUT}")


if __name__ == "__main__":
    main()
