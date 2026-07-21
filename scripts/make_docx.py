from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches


AUTHORS = (
    "Federico Lorenzo^1,2; Agustin Rosich^1,2; Jesica Lell^1,2; "
    "Sergio Aguiar^1; Valentina Ferreira^1; Karina Ochandorena^1,2; "
    "Eduardo Larrinaga^1; Natalia Gadea^1; Nicolas Larragueta^1; Aldo Quarneti^1"
)

AFFILIATIONS = (
    "^1 Radiotherapy, RT International Institute, Montevideo, Uruguay. "
    "^2 Radiotherapy, Unidad Academica de Radioterapia, Montevideo, Uruguay, RT International Institute."
)


def add_csv_table(document: Document, csv_path: Path) -> None:
    if not csv_path.exists():
        document.add_paragraph(f"[Missing table: {csv_path}]")
        return

    with csv_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))

    if not rows:
        return

    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for idx, header in enumerate(rows[0]):
        table.rows[0].cells[idx].text = header

    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row[: len(cells)]):
            cells[idx].text = value


def add_markdown_line(document: Document, line: str, base_dir: Path) -> None:
    stripped = line.strip()
    if not stripped:
        return

    if stripped.startswith("Authors:") or stripped.startswith("Affiliations:"):
        return

    heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if heading:
        level = min(len(heading.group(1)), 4)
        document.add_heading(heading.group(2), level=level)
        if len(heading.group(1)) == 1:
            document.add_paragraph(AUTHORS)
            document.add_paragraph(AFFILIATIONS)
        return

    if stripped.startswith("- "):
        document.add_paragraph(stripped[2:], style="List Bullet")
        return

    image_ref = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
    if image_ref:
        caption = image_ref.group(1)
        image_path = base_dir / image_ref.group(2)
        if image_path.exists():
            document.add_picture(str(image_path), width=Inches(6.2))
            if caption:
                document.add_paragraph(caption)
        else:
            document.add_paragraph(f"[Missing figure: {image_path}]")
        return

    table_ref = re.search(r"`(tables/[^`]+\.csv)`", stripped)
    if table_ref:
        document.add_paragraph(stripped.replace("`", ""))
        add_csv_table(document, base_dir / table_ref.group(1))
        return

    document.add_paragraph(stripped)


def build_docx(markdown_path: Path, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    text = markdown_path.read_text(encoding="utf-8")
    for line in text.splitlines():
        add_markdown_line(document, line, markdown_path.parent)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python scripts/make_docx.py manuscript.md output.docx")

    build_docx(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
